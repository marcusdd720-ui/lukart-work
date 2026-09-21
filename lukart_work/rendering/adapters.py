"""Stage 11 isolated DOCX/PDF adapters over normalized RenderTree."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from .contracts import RenderProfile, RenderTree


@dataclass(frozen=True)
class AdapterOutput:
    data: bytes
    media_type: str
    resource_hashes: tuple[str, ...]


def _styles(profile: RenderProfile):
    return {style.token: style for style in profile.styles}


def _normalize_docx_package(raw: bytes) -> bytes:
    source = BytesIO(raw)
    target = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(
        target, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = original.external_attr
            info.flag_bits = original.flag_bits & 0x800
            output.writestr(info, archive.read(name), compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()


def render_docx(tree: RenderTree, profile: RenderProfile) -> AdapterOutput:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    document = Document()
    fixed = datetime(2000, 1, 1, 0, 0, 0)
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    document.core_properties.author = "LUKART WORK"
    document.core_properties.last_modified_by = "LUKART WORK"
    document.core_properties.title = tree.semantic_digest

    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(profile.page.margin_top_mm)
        section.right_margin = Mm(profile.page.margin_right_mm)
        section.bottom_margin = Mm(profile.page.margin_bottom_mm)
        section.left_margin = Mm(profile.page.margin_left_mm)

    align = {
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
        "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    style_names = {}
    for index, spec in enumerate(profile.styles):
        name = f"LW_{index:02d}_{spec.token.replace('.', '_')}"
        style_names[spec.token] = name
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = profile.font.family
        style.font.size = Pt(spec.font_size_pt)
        style.font.bold = spec.bold
        style.font.italic = spec.italic
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), profile.font.family)
        rfonts.set(qn("w:hAnsi"), profile.font.family)
        rfonts.set(qn("w:eastAsia"), profile.font.family)
        fmt = style.paragraph_format
        fmt.alignment = align[spec.alignment]
        fmt.line_spacing = spec.line_spacing
        fmt.space_before = Pt(spec.space_before_pt)
        fmt.space_after = Pt(spec.space_after_pt)
        fmt.left_indent = Mm(spec.left_indent_mm)
        fmt.first_line_indent = Mm(spec.first_line_indent_mm)
        fmt.keep_with_next = spec.keep_with_next
        fmt.widow_control = True

    for block in tree.blocks:
        paragraph = document.add_paragraph(style=style_names[block.token])
        paragraph.add_run(block.text)

    raw = BytesIO()
    document.save(raw)
    data = _normalize_docx_package(raw.getvalue())
    return AdapterOutput(
        data=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        resource_hashes=(),
    )


_FONT_FILES = {
    "DejaVu Serif": {
        "R": "DejaVuSerif.ttf",
        "B": "DejaVuSerif-Bold.ttf",
        "I": "DejaVuSerif-Italic.ttf",
        "BI": "DejaVuSerif-BoldItalic.ttf",
    },
    "Liberation Serif": {
        "R": "LiberationSerif-Regular.ttf",
        "B": "LiberationSerif-Bold.ttf",
        "I": "LiberationSerif-Italic.ttf",
        "BI": "LiberationSerif-BoldItalic.ttf",
    },
}

_FONT_ROOTS = (
    "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/truetype/liberation2",
    "/usr/share/fonts/truetype/liberation",
    "/Library/Fonts",
    "C:/Windows/Fonts",
)


def _candidate_roots():
    explicit = os.environ.get("LUKART_FONT_DIR")
    roots = ([explicit] if explicit else []) + list(_FONT_ROOTS)
    return tuple(Path(root) for root in roots if root)


def _resolve_family(profile: RenderProfile):
    requested = (profile.font.family,) + (
        profile.font.fallback_families if profile.font.allow_fallback else ()
    )
    for family in requested:
        names = _FONT_FILES.get(family)
        if not names:
            continue
        resolved = {}
        for variant, filename in names.items():
            path = None
            for root in _candidate_roots():
                direct = root / filename
                if direct.is_file():
                    path = direct
                    break
            if path is None:
                resolved = {}
                break
            resolved[variant] = path
        if len(resolved) == 4:
            hashes = tuple(
                sha256(resolved[variant].read_bytes()).hexdigest()
                for variant in ("R", "B", "I", "BI")
            )
            return family, resolved, hashes
    raise ValueError("APPROVED_FONT_UNAVAILABLE")


def _pdf_fonts(profile: RenderProfile):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    family, paths, hashes = _resolve_family(profile)
    family_digest = sha256(("|".join(hashes)).encode("ascii")).hexdigest()[:16]
    aliases = {
        "R": f"LW{family_digest}R",
        "B": f"LW{family_digest}B",
        "I": f"LW{family_digest}I",
        "BI": f"LW{family_digest}BI",
    }
    registered = set(pdfmetrics.getRegisteredFontNames())
    for variant in ("R", "B", "I", "BI"):
        if aliases[variant] not in registered:
            pdfmetrics.registerFont(TTFont(aliases[variant], str(paths[variant])))
    return family, aliases, hashes


def _wrap(text: str, font_name: str, size: float, width: float):
    from reportlab.pdfbase import pdfmetrics

    words = text.split()
    if not words:
        return [""]
    lines = []
    current = []
    for word in words:
        candidate = " ".join((*current, word))
        if pdfmetrics.stringWidth(candidate, font_name, size) <= width:
            current.append(word)
            continue
        if current:
            lines.append(" ".join(current))
            current = []
        if pdfmetrics.stringWidth(word, font_name, size) <= width:
            current = [word]
            continue
        chunk = ""
        for char in word:
            if chunk and pdfmetrics.stringWidth(chunk + char, font_name, size) > width:
                lines.append(chunk)
                chunk = char
            else:
                chunk += char
        current = [chunk] if chunk else []
    if current:
        lines.append(" ".join(current))
    return lines or [""]


def _font_variant(style):
    if style.bold and style.italic:
        return "BI"
    if style.bold:
        return "B"
    if style.italic:
        return "I"
    return "R"


def render_pdf(tree: RenderTree, profile: RenderProfile) -> AdapterOutput:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas

    _, aliases, resource_hashes = _pdf_fonts(profile)
    style_map = _styles(profile)
    mm = 72.0 / 25.4
    page_width, page_height = A4
    left = profile.page.margin_left_mm * mm
    right = profile.page.margin_right_mm * mm
    top = profile.page.margin_top_mm * mm
    bottom = profile.page.margin_bottom_mm * mm
    usable_width = page_width - left - right

    output = BytesIO()
    pdf = canvas.Canvas(
        output,
        pagesize=A4,
        invariant=1,
        pageCompression=0,
        bottomup=1,
    )
    pdf.setAuthor("LUKART WORK")
    pdf.setCreator("LUKART WORK Stage 11")
    pdf.setTitle(tree.semantic_digest)
    pdf.setSubject(tree.profile_digest)

    y = page_height - top

    def new_page():
        nonlocal y
        pdf.showPage()
        y = page_height - top

    for block_index, block in enumerate(tree.blocks):
        style = style_map[block.token]
        font_name = aliases[_font_variant(style)]
        size = style.font_size_pt
        line_height = size * style.line_spacing
        width = usable_width - style.left_indent_mm * mm
        lines = _wrap(block.text, font_name, size, width)
        needed = style.space_before_pt + len(lines) * line_height + style.space_after_pt
        if style.keep_with_next and block_index + 1 < len(tree.blocks):
            needed += line_height * 2
        if y - needed < bottom:
            new_page()
        y -= style.space_before_pt
        x0 = left + style.left_indent_mm * mm

        for line_index, line in enumerate(lines):
            is_last = line_index == len(lines) - 1
            if style.alignment == "CENTER":
                x = x0 + max(0, width - pdfmetrics.stringWidth(line, font_name, size)) / 2
                pdf.setFont(font_name, size)
                pdf.drawString(x, y - size, line)
            elif style.alignment == "RIGHT":
                x = x0 + max(0, width - pdfmetrics.stringWidth(line, font_name, size))
                pdf.setFont(font_name, size)
                pdf.drawString(x, y - size, line)
            elif style.alignment == "JUSTIFY" and not is_last and len(line.split()) > 1:
                words = line.split()
                word_width = sum(pdfmetrics.stringWidth(word, font_name, size) for word in words)
                gap = max(0, (width - word_width) / (len(words) - 1))
                x = x0
                pdf.setFont(font_name, size)
                for word in words:
                    pdf.drawString(x, y - size, word)
                    x += pdfmetrics.stringWidth(word, font_name, size) + gap
            else:
                x = x0 + (style.first_line_indent_mm * mm if line_index == 0 else 0)
                pdf.setFont(font_name, size)
                pdf.drawString(x, y - size, line)
            y -= line_height
        y -= style.space_after_pt

    pdf.save()
    return AdapterOutput(
        data=output.getvalue(),
        media_type="application/pdf",
        resource_hashes=resource_hashes,
    )


def render_format(output_format: str, tree: RenderTree, profile: RenderProfile) -> AdapterOutput:
    if output_format == "DOCX":
        return render_docx(tree, profile)
    if output_format == "PDF":
        return render_pdf(tree, profile)
    raise ValueError("UNSUPPORTED_FORMAT")

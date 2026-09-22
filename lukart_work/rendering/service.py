"""Deterministic Stage 11 Rendering Engine.

The renderer is deliberately non-generative: it resolves already-approved semantic data,
builds a renderer-neutral model, lays it out, emits DOCX/PDF and validates the resulting
artifacts.  It never changes legal content or case state.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from importlib.metadata import version as metadata_version
from pathlib import Path
from typing import Mapping
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab import pdfbase
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from lukart_work.contracts.messages import RenderRequest, SemanticDocument, ValidationRunPin
from lukart_work.domain.models import Decision, DocumentPlan, Draft, Entity, Evidence, Institution, Party, Ref, Request, Section
from lukart_work.validation.contracts import PRE_RENDER_CHECKS, ValidationProjection
from lukart_work.validation.policy import validation_pin
from lukart_work.validation.rules import Violation, digest, ref as entity_ref

from .contracts import (
    LayoutBlock,
    LayoutTree,
    RenderArtifact,
    RenderAuthority,
    RenderBlock,
    RenderBundleMeta,
    RenderModel,
    RenderProfile,
    RenderReceipt,
)
from .profiles import OFFICIAL_PL_V1, profile_font_paths, profile_pin, profile_styles


STYLE_TOKEN = {
    "DOCUMENT_HEADER": "document.heading.primary",
    "PARTY_BLOCK": "document.party_block",
    "REQUEST_HEADING": "document.heading.secondary",
    "REQUEST_ITEM": "document.request",
    "SECTION_HEADING": "document.heading.secondary",
    "BODY": "document.body",
    "LEGAL_BASIS": "document.legal_basis",
    "QUOTE": "document.quote",
    "SIGNATURE": "document.signature",
    "ATTACHMENT_HEADING": "document.heading.secondary",
    "ATTACHMENT_ITEM": "document.attachment",
}

FIXED_DOC_TIME = datetime(2000, 1, 1, 0, 0, 0)
FIXED_ZIP_TIME = (2000, 1, 1, 0, 0, 0)
RENDERER_VERSION = "1.0.0"


def _sha(data: bytes) -> str:
    return sha256(data).hexdigest()


def authorization_binding(
    document: SemanticDocument,
    profile: "VersionPin",
    authority: RenderAuthority,
) -> str:
    """Bind exact semantic input, render profile and authoritative Stage 10 evidence.

    This is an integrity binding, not an authentication mechanism. The trusted
    pre-render boundary remains responsible for deciding that the exact bound input
    is authorized for rendering.
    """
    return digest({
        "semantic_digest": digest(document),
        "render_profile": profile.model_dump(mode="json"),
        "authority": authority.model_dump(mode="json"),
    })


def renderer_pin() -> "VersionPin":
    from lukart_work.domain.models import VersionPin

    root = Path(__file__).resolve().parent
    names = ("contracts.py", "profiles.py", "service.py")
    identity = {
        "sources": {
            name: _sha((root / name).read_bytes().replace(b"\r\n", b"\n"))
            for name in names
        },
        "dependencies": {
            "python-docx": metadata_version("python-docx"),
            "reportlab": metadata_version("reportlab"),
            "fontpkg": metadata_version("fontpkg"),
            "fontpkg-lato": metadata_version("fontpkg-lato"),
        },
    }
    return VersionPin(
        id="lukart-work-renderer",
        version=RENDERER_VERSION,
        content_hash=digest(identity),
    )


def _record(records: Mapping[Ref, object], target: Ref, scope: tuple[object, object]):
    item = records.get(target)
    if item is None:
        raise Violation("BROKEN_REFERENCE")
    if not isinstance(item, Entity) or entity_ref(item) != target:
        raise Violation("INVALID_HASH")
    if (item.tenant_id, item.case_id) != scope:
        raise Violation("CROSS_CASE")
    return item


def derive_render_authority(
    document: SemanticDocument,
    records: Mapping[Ref, object],
    validation: ValidationProjection,
    render_profile: "VersionPin",
) -> RenderAuthority:
    """Verify and pin the exact current Stage 10 authority used by the renderer."""
    scope = (document.tenant_id, document.case_id)
    execution = validation.execution
    run = execution.run

    if validation.freshness != "CURRENT" or not validation.eligible_for_approval or validation.reasons:
        raise Violation("STALE")
    if execution.id != run.id:
        raise Violation("INVALID_HASH")

    current_policy = validation_pin()
    if (
        execution.validator != current_policy
        or run.policy != current_policy
        or execution.manifest.policy != current_policy
    ):
        raise Violation("STALE")
    if (
        digest(execution.manifest) != run.execution_manifest_hash
        or execution.manifest.draft_hash != run.draft_hash
        or execution.manifest.snapshot_hash != run.snapshot_hash
        or execution.manifest.review_hash != (
            digest(execution.review) if execution.review is not None else None
        )
    ):
        raise Violation("INVALID_HASH")

    input_draft = _record(records, run.draft, scope)
    if not isinstance(input_draft, Draft) or digest(input_draft) != run.draft_hash:
        raise Violation("INVALID_HASH")

    output_draft = _record(records, execution.output_draft, scope)
    if not isinstance(output_draft, Draft):
        raise Violation("REFERENCE_TYPE")
    if (
        entity_ref(output_draft) != document.draft
        or digest(output_draft) != execution.output_hash
        or output_draft.status != "VALIDATED"
    ):
        raise Violation("STALE")
    if (
        output_draft.snapshot_hash != document.snapshot_hash
        or run.snapshot_hash != output_draft.snapshot_hash
        or execution.manifest.snapshot_hash != output_draft.snapshot_hash
    ):
        raise Violation("INVALID_HASH")
    if output_draft.reproduction.render_profile != render_profile:
        raise Violation("GATE_DENIED")
    if document.sections != output_draft.sections:
        raise Violation("INVALID_HASH")

    dependency_pairs = execution.manifest.dependency_hashes
    dependency_refs = tuple(target for target, _ in dependency_pairs)
    if dependency_refs != execution.dependencies or len(set(dependency_refs)) != len(dependency_refs):
        raise Violation("INVALID_HASH")
    for target, expected_hash in dependency_pairs:
        item = _record(records, target, scope)
        if digest(item) != expected_hash:
            raise Violation("STALE")

    plan = _record(records, output_draft.plan, scope)
    if not isinstance(plan, DocumentPlan):
        raise Violation("REFERENCE_TYPE")
    if (
        plan.status != "APPROVED"
        or plan.approval is None
        or plan.input_snapshot_id != output_draft.snapshot_id
        or plan.input_snapshot_hash != output_draft.snapshot_hash
        or tuple(section.section for section in output_draft.sections) != plan.section_order
    ):
        raise Violation("GATE_DENIED")
    if output_draft.plan not in dependency_refs or plan.approval not in dependency_refs:
        raise Violation("STALE")
    approval = _record(records, plan.approval, scope)
    if not isinstance(approval, Decision):
        raise Violation("REFERENCE_TYPE")
    if (
        approval.status != "RECORDED"
        or approval.decision_type != "PLAN_APPROVAL"
        or approval.outcome != "ACCEPT"
        or entity_ref(plan) not in approval.targets
    ):
        raise Violation("GATE_DENIED")

    checks = {check.check: check for check in run.checks}
    expected_checks = set(PRE_RENDER_CHECKS) | {"render"}
    if len(checks) != len(run.checks) or set(checks) != expected_checks:
        raise Violation("GATE_DENIED")
    if any(checks[name].outcome != "PASS" for name in PRE_RENDER_CHECKS):
        raise Violation("GATE_DENIED")
    if checks["render"].outcome != "NOT_RUN":
        raise Violation("GATE_DENIED")

    return RenderAuthority(
        semantic_digest=digest(document),
        draft=entity_ref(output_draft),
        draft_hash=digest(output_draft),
        plan=entity_ref(plan),
        plan_hash=digest(plan),
        snapshot_id=output_draft.snapshot_id,
        snapshot_hash=output_draft.snapshot_hash,
        validation_run=ValidationRunPin(id=run.id, content_hash=digest(run)),
        validation_execution_hash=digest(execution),
        validation_projection_hash=digest(validation),
        case_sequence=execution.committed_sequence,
    )


def _party_text(item: object) -> str:
    if isinstance(item, Party):
        name = item.display_name
    elif isinstance(item, Institution):
        name = item.name
    else:
        raise Violation("REFERENCE_TYPE")
    address = getattr(item, "address", None)
    if address is None:
        return name
    lines = [name, *address.lines]
    if address.postal_code:
        lines.append(address.postal_code)
    return "\n".join(lines)


def _attachment_label(records: Mapping[Ref, object], anchor, scope) -> tuple[str, Ref]:
    matches = [
        (target, item)
        for target, item in records.items()
        if isinstance(item, Evidence)
        and (item.tenant_id, item.case_id) == scope
        and item.source == anchor
    ]
    if len(matches) != 1:
        raise Violation("PROVENANCE_MISSING")
    target, evidence = matches[0]
    return evidence.description, target


def assemble_render_model(
    document: SemanticDocument,
    records: Mapping[Ref, object],
    authority: RenderAuthority,
    profile: RenderProfile = OFFICIAL_PL_V1,
) -> RenderModel:
    scope = (document.tenant_id, document.case_id)
    draft = _record(records, document.draft, scope)
    if not isinstance(draft, Draft):
        raise Violation("REFERENCE_TYPE")
    if (
        draft.snapshot_hash != document.snapshot_hash
        or document.sections != draft.sections
        or authority.semantic_digest != digest(document)
        or authority.draft != document.draft
        or authority.draft_hash != digest(draft)
        or authority.plan != draft.plan
        or authority.snapshot_id != draft.snapshot_id
        or authority.snapshot_hash != draft.snapshot_hash
    ):
        raise Violation("INVALID_HASH")
    if draft.status != "VALIDATED":
        raise Violation("GATE_DENIED")

    blocks: list[RenderBlock] = [
        RenderBlock(role="DOCUMENT_HEADER", text=document.header, refs=(document.draft,))
    ]

    for party_ref in document.parties:
        party = _record(records, party_ref, scope)
        blocks.append(RenderBlock(role="PARTY_BLOCK", text=_party_text(party), refs=(party_ref,)))

    if document.requests:
        blocks.append(RenderBlock(role="REQUEST_HEADING", text=profile.requests_label))
        for ordinal, request_ref in enumerate(document.requests, 1):
            request = _record(records, request_ref, scope)
            if not isinstance(request, Request):
                raise Violation("REFERENCE_TYPE")
            blocks.append(
                RenderBlock(
                    role="REQUEST_ITEM",
                    text=request.text,
                    refs=(request_ref,),
                    ordinal=ordinal,
                )
            )

    statement_roles = {
        "FACTUAL": "BODY",
        "LEGAL": "LEGAL_BASIS",
        "REQUEST": "REQUEST_ITEM",
        "INFERENCE": "BODY",
        "STRUCTURAL": "BODY",
    }
    for section_text in document.sections:
        section = _record(records, section_text.section, scope)
        if not isinstance(section, Section):
            raise Violation("REFERENCE_TYPE")
        blocks.append(
            RenderBlock(role="SECTION_HEADING", text=section.title, refs=(section_text.section,))
        )
        request_ordinal = 1
        for statement in section_text.statements:
            role = statement_roles[statement.statement_type]
            ordinal = None
            if role == "REQUEST_ITEM":
                ordinal = request_ordinal
                request_ordinal += 1
            refs = tuple(dict.fromkeys((section_text.section, *statement.refs, *statement.citations)))
            blocks.append(
                RenderBlock(role=role, text=statement.text, refs=refs, ordinal=ordinal)
            )

    if document.attachments:
        blocks.append(RenderBlock(role="ATTACHMENT_HEADING", text=profile.attachments_label))
        for ordinal, anchor in enumerate(document.attachments, 1):
            label, evidence_ref = _attachment_label(records, anchor, scope)
            blocks.append(
                RenderBlock(
                    role="ATTACHMENT_ITEM",
                    text=label,
                    refs=(evidence_ref,),
                    ordinal=ordinal,
                )
            )

    if document.signature_party is not None:
        party = _record(records, document.signature_party, scope)
        blocks.append(
            RenderBlock(
                role="SIGNATURE",
                text=_party_text(party).splitlines()[0],
                refs=(document.signature_party,),
            )
        )

    return RenderModel(
        tenant_id=document.tenant_id,
        case_id=document.case_id,
        draft=document.draft,
        snapshot_hash=document.snapshot_hash,
        authority=authority,
        blocks=tuple(blocks),
    )


def build_layout(model: RenderModel, profile: RenderProfile = OFFICIAL_PL_V1) -> LayoutTree:
    blocks = []
    for block in model.blocks:
        token = STYLE_TOKEN.get(block.role)
        if token is None:
            raise Violation("SCHEMA_INVALID")
        blocks.append(
            LayoutBlock(
                role=block.role,
                style_token=token,
                text=block.text,
                refs=block.refs,
                ordinal=block.ordinal,
            )
        )
    return LayoutTree(
        tenant_id=model.tenant_id,
        case_id=model.case_id,
        draft=model.draft,
        snapshot_hash=model.snapshot_hash,
        authority=model.authority,
        profile=profile_pin(profile),
        blocks=tuple(blocks),
    )


def _display_text(block: LayoutBlock) -> str:
    return f"{block.ordinal}. {block.text}" if block.ordinal is not None else block.text


def _set_run_font(run, profile: RenderProfile, *, size: int, bold: bool) -> None:
    run.font.name = profile.font_family
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), profile.font_family)


def _normalize_docx(raw: bytes) -> bytes:
    source = BytesIO(raw)
    output = BytesIO()
    with ZipFile(source, "r") as src, ZipFile(
        output, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as dst:
        for name in sorted(src.namelist()):
            info = ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            dst.writestr(info, src.read(name))
    return output.getvalue()


def render_docx(layout: LayoutTree, profile: RenderProfile = OFFICIAL_PL_V1) -> bytes:
    if layout.profile != profile_pin(profile):
        raise Violation("INVALID_HASH")
    document = Document()
    core = document.core_properties
    core.author = "LUKART WORK"
    core.last_modified_by = "LUKART WORK"
    core.title = ""
    core.subject = ""
    core.keywords = ""
    core.comments = ""
    core.created = FIXED_DOC_TIME
    core.modified = FIXED_DOC_TIME
    core.revision = 1

    section = document.sections[0]
    section.page_width = Mm(profile.page_width_mm)
    section.page_height = Mm(profile.page_height_mm)
    section.top_margin = Mm(profile.margin_top_mm)
    section.right_margin = Mm(profile.margin_right_mm)
    section.bottom_margin = Mm(profile.margin_bottom_mm)
    section.left_margin = Mm(profile.margin_left_mm)

    if profile.header_text:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run(profile.header_text)
        _set_run_font(run, profile, size=profile.header_footer_font_pt, bold=False)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if profile.footer_text:
        run = footer.add_run(profile.footer_text)
        _set_run_font(run, profile, size=profile.header_footer_font_pt, bold=False)
    if profile.page_numbers:
        if profile.footer_text:
            separator = footer.add_run(" — ")
            _set_run_font(separator, profile, size=profile.header_footer_font_pt, bold=False)
        page_run = footer.add_run()
        _set_run_font(page_run, profile, size=profile.header_footer_font_pt, bold=False)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        display = OxmlElement("w:t")
        display.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        page_run._r.extend((begin, instruction, separate, display, end))

    # Remove the package's initial empty paragraph so artifact identity depends only on layout.
    body = document._element.body
    for paragraph in list(body.findall(qn("w:p"))):
        body.remove(paragraph)

    styles = profile_styles(profile)
    alignments = {
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
        "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    for block in layout.blocks:
        style = styles.get(block.style_token)
        if style is None:
            raise Violation("RENDER_STYLE_UNSUPPORTED")
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.line_spacing = style.line_spacing_milli / 1000
        fmt.space_before = Pt(style.space_before_pt)
        fmt.space_after = Pt(style.space_after_pt)
        fmt.left_indent = Mm(style.left_indent_mm)
        fmt.first_line_indent = Mm(-style.hanging_indent_mm) if style.hanging_indent_mm else None
        fmt.keep_with_next = style.keep_with_next
        fmt.keep_together = style.keep_together
        fmt.widow_control = style.widow_control
        fmt.page_break_before = style.page_break_before
        paragraph.alignment = alignments[style.alignment]

        run = paragraph.add_run(_display_text(block))
        _set_run_font(run, profile, size=style.font_size_pt, bold=style.bold)

    buffer = BytesIO()
    document.save(buffer)
    return _normalize_docx(buffer.getvalue())


def _register_pdf_fonts(profile: RenderProfile) -> tuple[str, str]:
    regular, bold = profile_font_paths(profile)
    regular_name = f"LUKART-{profile.profile_id}-{profile.version}-Regular"
    bold_name = f"LUKART-{profile.profile_id}-{profile.version}-Bold"
    if regular_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(regular_name, str(regular)))
    if bold_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(bold_name, str(bold)))
    return regular_name, bold_name


def _wrap_pdf(text: str, font_name: str, size: int, max_width: float) -> list[str]:
    lines: list[str] = []
    raw_lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for raw in raw_lines:
        if raw == "":
            lines.append("")
            continue
        words = raw.split(" ")
        current = ""
        for word in words:
            candidate = word if not current else f"{current} {word}"
            if not current or pdfmetrics.stringWidth(candidate, font_name, size) <= max_width:
                current = candidate
                continue
            lines.append(current)
            current = word
            if pdfmetrics.stringWidth(current, font_name, size) > max_width:
                # A single pathological token must not escape the page box.
                piece = ""
                for char in current:
                    test = piece + char
                    if piece and pdfmetrics.stringWidth(test, font_name, size) > max_width:
                        lines.append(piece)
                        piece = char
                    else:
                        piece = test
                current = piece
        lines.append(current)
    return lines


def render_pdf(layout: LayoutTree, profile: RenderProfile = OFFICIAL_PL_V1) -> bytes:
    if layout.profile != profile_pin(profile):
        raise Violation("INVALID_HASH")
    regular_name, bold_name = _register_pdf_fonts(profile)
    page_width = profile.page_width_mm * mm
    page_height = profile.page_height_mm * mm
    left = profile.margin_left_mm * mm
    right = page_width - profile.margin_right_mm * mm
    top = page_height - profile.margin_top_mm * mm
    bottom = profile.margin_bottom_mm * mm
    width = right - left

    output = BytesIO()
    pdf = canvas.Canvas(
        output,
        pagesize=(page_width, page_height),
        invariant=1,
        pageCompression=1,
    )
    pdf.setAuthor("LUKART WORK")
    pdf.setCreator(f"LUKART WORK Renderer {RENDERER_VERSION}")
    pdf.setTitle("")
    pdf.setSubject("")

    y = top
    page_number = 1

    def draw_running_elements() -> None:
        pdf.setFont(regular_name, profile.header_footer_font_pt)
        if profile.header_text:
            pdf.drawCentredString(page_width / 2, page_height - profile.footer_offset_mm * mm, profile.header_text)
        footer_parts = []
        if profile.footer_text:
            footer_parts.append(profile.footer_text)
        if profile.page_numbers:
            footer_parts.append(str(page_number))
        if footer_parts:
            pdf.drawCentredString(
                page_width / 2,
                profile.footer_offset_mm * mm,
                " — ".join(footer_parts),
            )

    def new_page() -> None:
        nonlocal y, page_number
        draw_running_elements()
        pdf.showPage()
        page_number += 1
        y = top

    styles = profile_styles(profile)
    for block_index, block in enumerate(layout.blocks):
        style = styles.get(block.style_token)
        if style is None:
            raise Violation("RENDER_STYLE_UNSUPPORTED")
        size = style.font_size_pt
        font = bold_name if style.bold else regular_name
        before = float(style.space_before_pt)
        after = float(style.space_after_pt)
        line_height = size * (style.line_spacing_milli / 1000)
        text = _display_text(block)
        block_left = left + style.left_indent_mm * mm
        available = right - block_left
        lines = _wrap_pdf(text, font, size, available)

        if style.page_break_before and y != top:
            new_page()

        reserve = before + max(1, len(lines)) * line_height + after
        if style.keep_with_next and block_index + 1 < len(layout.blocks):
            next_block = layout.blocks[block_index + 1]
            next_style = styles.get(next_block.style_token)
            if next_style is None:
                raise Violation("RENDER_STYLE_UNSUPPORTED")
            reserve += next_style.space_before_pt + (
                next_style.font_size_pt * (next_style.line_spacing_milli / 1000)
            )
        if style.keep_together and reserve <= (top - bottom) and y - reserve < bottom:
            new_page()
        elif style.keep_with_next and y - reserve < bottom:
            new_page()

        y -= before
        pdf.setFont(font, size)
        for index, line in enumerate(lines):
            remaining = len(lines) - index
            capacity = int(max(0, y - bottom) // line_height)
            if y - line_height < bottom or (
                style.widow_control and remaining > 1 and capacity == 1
            ):
                new_page()
                pdf.setFont(font, size)
            if style.alignment == "CENTER":
                pdf.drawCentredString(page_width / 2, y, line)
            elif style.alignment == "RIGHT":
                pdf.drawRightString(right, y, line)
            elif style.alignment == "JUSTIFY" and index < len(lines) - 1 and " " in line:
                words = line.split()
                if len(words) > 1:
                    text_width = sum(pdfmetrics.stringWidth(word, font, size) for word in words)
                    word_space = max(0.0, (available - text_width) / (len(words) - 1))
                    cursor = block_left
                    for word in words:
                        pdf.drawString(cursor, y, word)
                        cursor += pdfmetrics.stringWidth(word, font, size) + word_space
                else:
                    pdf.drawString(block_left, y, line)
            else:
                pdf.drawString(block_left, y, line)
            y -= line_height
        y -= after

    draw_running_elements()
    pdf.save()
    return output.getvalue()


def _validate_docx(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data), "r") as package:
            names = set(package.namelist())
            required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
            if not required <= names:
                raise Violation("RENDER_INVALID")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise Violation("RENDER_INVALID")
            for name in names:
                if name.endswith(".rels") and b'TargetMode="External"' in package.read(name):
                    raise Violation("RENDER_INVALID")
    except Violation:
        raise
    except Exception as exc:
        raise Violation("RENDER_INVALID") from exc


def _validate_pdf(data: bytes) -> None:
    if not data.startswith(b"%PDF-"):
        raise Violation("RENDER_INVALID")
    for forbidden in (b"/JavaScript", b"/JS ", b"/OpenAction", b"/Launch"):
        if forbidden in data:
            raise Violation("RENDER_INVALID")
    if b"/FontFile2" not in data:
        raise Violation("RENDER_FONT_MISSING")


@dataclass(frozen=True)
class RenderedBundle:
    docx: bytes
    pdf: bytes
    meta: RenderBundleMeta


class RenderingEngine:
    def render(
        self,
        request: RenderRequest,
        records: Mapping[Ref, object],
        validation: ValidationProjection,
        profile: RenderProfile = OFFICIAL_PL_V1,
    ) -> RenderedBundle:
        if request.profile != profile_pin(profile):
            raise Violation("INVALID_HASH")
        authority = derive_render_authority(
            request.document, records, validation, request.profile
        )
        if request.authorization_hash != authorization_binding(
            request.document, request.profile, authority
        ):
            raise Violation("GATE_DENIED")
        model = assemble_render_model(request.document, records, authority, profile)
        layout = build_layout(model, profile)
        docx = render_docx(layout, profile)
        pdf = render_pdf(layout, profile)
        _validate_docx(docx)
        _validate_pdf(pdf)

        artifacts = (
            RenderArtifact(
                format="DOCX",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                content_hash=_sha(docx),
                byte_length=len(docx),
            ),
            RenderArtifact(
                format="PDF",
                media_type="application/pdf",
                content_hash=_sha(pdf),
                byte_length=len(pdf),
            ),
        )
        receipt = RenderReceipt(
            draft=model.draft,
            snapshot_hash=model.snapshot_hash,
            authority=model.authority,
            authorization_hash=request.authorization_hash,
            semantic_digest=digest(request.document),
            render_model_digest=digest(model),
            layout_digest=digest(layout),
            render_profile=profile_pin(profile),
            renderer=renderer_pin(),
            artifacts=artifacts,
        )
        meta = RenderBundleMeta(model=model, layout=layout, receipt=receipt)
        return RenderedBundle(docx=docx, pdf=pdf, meta=meta)

    @staticmethod
    def validate(bundle: RenderedBundle) -> None:
        _validate_docx(bundle.docx)
        _validate_pdf(bundle.pdf)
        artifacts = bundle.meta.receipt.artifacts
        if (
            artifacts[0].content_hash != _sha(bundle.docx)
            or artifacts[0].byte_length != len(bundle.docx)
            or artifacts[1].content_hash != _sha(bundle.pdf)
            or artifacts[1].byte_length != len(bundle.pdf)
        ):
            raise Violation("INVALID_HASH")
        if bundle.meta.receipt.render_model_digest != digest(bundle.meta.model):
            raise Violation("INVALID_HASH")
        if bundle.meta.receipt.layout_digest != digest(bundle.meta.layout):
            raise Violation("INVALID_HASH")
        if not (
            bundle.meta.receipt.authority == bundle.meta.model.authority
            == bundle.meta.layout.authority
        ):
            raise Violation("INVALID_HASH")
        if bundle.meta.receipt.authority.semantic_digest != bundle.meta.receipt.semantic_digest:
            raise Violation("INVALID_HASH")
